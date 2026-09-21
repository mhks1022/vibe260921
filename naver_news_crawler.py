"""
네이버 통합검색 결과에서 뉴스 기사 링크를 찾아
각 기사의 제목/언론사/날짜/본문을 크롤링하는 PyQt6 GUI 앱

동작 방식
---------
네이버 통합검색(nexearch) 결과 페이지는 대부분의 콘텐츠가 자바스크립트로
렌더링되며 마크업 클래스명도 자주 바뀌기 때문에, 검색 결과 페이지 자체를
CSS 선택자로 파싱하는 방식은 쉽게 깨진다. 대신 이 스크립트는:

1. 검색 결과 페이지의 원본 HTML에서 정규식으로 네이버뉴스
   (n.news.naver.com) 기사 링크만 추출하고,
2. 각 기사 링크에 실제로 접속해 안정적인 네이버뉴스 기사 템플릿
   (#title_area, #dic_area 등)에서 제목/언론사/날짜/본문을 가져온다.

크롤링은 UI가 멈추지 않도록 QThread에서 실행하며, 결과는 표(테이블)로
보여주고 선택한 기사의 본문을 미리보기로 확인한 뒤 CSV로 저장할 수 있다.
"""

import csv
import re
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from PyQt6.QtCore import QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

SEARCH_URL = (
    "https://search.naver.com/search.naver"
    "?where=nexearch&sm=top_hty&fbm=0&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4"
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
}

REQUEST_DELAY = 1.0  # 초 단위, 요청 사이 대기 시간 (서버 부하 방지)

ARTICLE_LINK_PATTERN = re.compile(r"n\.news\.naver\.com/mnews/article/(\d+)/(\d+)")


def get_soup(url):
    res = requests.get(url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = "utf-8"
    return BeautifulSoup(res.text, "html.parser")


def fetch_article_links(search_url, max_links=10):
    """검색 결과 페이지 원본 HTML에서 네이버뉴스 기사 링크를 정규식으로 추출한다."""
    res = requests.get(search_url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = "utf-8"

    links = []
    seen = set()
    for oid, aid in ARTICLE_LINK_PATTERN.findall(res.text):
        url = f"https://n.news.naver.com/mnews/article/{oid}/{aid}"
        if url not in seen:
            seen.add(url)
            links.append(url)
        if len(links) >= max_links:
            break

    return links


def fetch_article(link):
    """네이버뉴스 기사 페이지에서 제목/언론사/날짜/본문을 추출한다."""
    try:
        soup = get_soup(link)
    except requests.RequestException as e:
        return {"title": "", "press": "", "date": "", "content": f"[본문 요청 실패: {e}]"}

    title_tag = soup.select_one("#title_area")
    press_tag = soup.select_one(".media_end_head_top a img")
    date_tag = soup.select_one(".media_end_head_info_datestamp_time")
    body_tag = soup.select_one("#dic_area") or soup.select_one("#articleBodyContents")

    return {
        "title": title_tag.get_text(strip=True) if title_tag else "",
        "press": press_tag.get("alt", "") if press_tag else "",
        "date": date_tag.get("data-date-time", date_tag.get_text(strip=True)) if date_tag else "",
        "content": body_tag.get_text("\n", strip=True) if body_tag else "[본문을 찾을 수 없습니다]",
    }


RESULT_FIELDS = ["title", "press", "date", "link", "content"]
RESULT_HEADERS = ["제목", "언론사", "날짜", "링크", "본문"]


def save_to_csv(results, csv_path="naver_news_result.csv"):
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=RESULT_FIELDS)
        writer.writeheader()
        writer.writerows(results)


def save_to_excel(results, xlsx_path="naver_news_result.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "네이버뉴스"

    ws.append(RESULT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for article in results:
        ws.append([article.get(field, "") for field in RESULT_FIELDS])

    column_widths = [40, 14, 20, 45, 60]
    for col_idx, width in enumerate(column_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    wb.save(xlsx_path)


class CrawlWorker(QThread):
    progress = pyqtSignal(int, int, dict)
    finished_all = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, search_url, max_articles):
        super().__init__()
        self.search_url = search_url
        self.max_articles = max_articles

    def run(self):
        try:
            links = fetch_article_links(self.search_url, max_links=self.max_articles)
        except requests.RequestException as e:
            self.error.emit(f"검색 결과를 가져오지 못했습니다:\n{e}")
            return

        results = []
        total = len(links)
        for idx, link in enumerate(links, start=1):
            article = fetch_article(link)
            article["link"] = link
            results.append(article)
            self.progress.emit(idx, total, article)
            time.sleep(REQUEST_DELAY)

        self.finished_all.emit(results)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("네이버 뉴스 크롤러")
        self.resize(920, 680)

        self.results = []
        self.worker = None

        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)

        top_layout = QHBoxLayout()
        self.url_input = QLineEdit(SEARCH_URL)
        self.count_input = QSpinBox()
        self.count_input.setRange(1, 50)
        self.count_input.setValue(10)
        self.start_btn = QPushButton("크롤링 시작")
        self.start_btn.clicked.connect(self.start_crawl)

        top_layout.addWidget(QLabel("검색 URL:"))
        top_layout.addWidget(self.url_input, 1)
        top_layout.addWidget(QLabel("최대 기사 수:"))
        top_layout.addWidget(self.count_input)
        top_layout.addWidget(self.start_btn)
        layout.addLayout(top_layout)

        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["번호", "제목", "언론사", "날짜"])
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.itemSelectionChanged.connect(self.show_selected_content)
        layout.addWidget(self.table, 2)

        layout.addWidget(QLabel("본문 미리보기:"))
        self.content_view = QTextEdit()
        self.content_view.setReadOnly(True)
        layout.addWidget(self.content_view, 3)

        bottom_layout = QHBoxLayout()
        self.status_label = QLabel("대기 중")
        self.save_csv_btn = QPushButton("CSV로 저장")
        self.save_csv_btn.setEnabled(False)
        self.save_csv_btn.clicked.connect(self.save_csv)
        self.save_excel_btn = QPushButton("엑셀로 저장")
        self.save_excel_btn.setEnabled(False)
        self.save_excel_btn.clicked.connect(self.save_excel)
        bottom_layout.addWidget(self.status_label, 1)
        bottom_layout.addWidget(self.save_csv_btn)
        bottom_layout.addWidget(self.save_excel_btn)
        layout.addLayout(bottom_layout)

    def start_crawl(self):
        search_url = self.url_input.text().strip()
        if not search_url:
            QMessageBox.warning(self, "입력 오류", "검색 URL을 입력해주세요.")
            return

        self.table.setRowCount(0)
        self.content_view.clear()
        self.results = []
        self.save_csv_btn.setEnabled(False)
        self.save_excel_btn.setEnabled(False)
        self.start_btn.setEnabled(False)
        self.status_label.setText("기사 링크를 찾는 중...")

        self.worker = CrawlWorker(search_url, self.count_input.value())
        self.worker.progress.connect(self.on_progress)
        self.worker.finished_all.connect(self.on_finished)
        self.worker.error.connect(self.on_error)
        self.worker.start()

    def on_progress(self, idx, total, article):
        self.results.append(article)
        row = self.table.rowCount()
        self.table.insertRow(row)
        self.table.setItem(row, 0, QTableWidgetItem(str(idx)))
        self.table.setItem(row, 1, QTableWidgetItem(article["title"]))
        self.table.setItem(row, 2, QTableWidgetItem(article["press"]))
        self.table.setItem(row, 3, QTableWidgetItem(article["date"]))
        self.status_label.setText(f"크롤링 중... ({idx}/{total})")

    def on_finished(self, results):
        self.results = results
        self.status_label.setText(f"완료: 총 {len(results)}건")
        self.start_btn.setEnabled(True)
        self.save_csv_btn.setEnabled(len(results) > 0)
        self.save_excel_btn.setEnabled(len(results) > 0)

    def on_error(self, message):
        self.status_label.setText("오류 발생")
        self.start_btn.setEnabled(True)
        QMessageBox.critical(self, "오류", message)

    def show_selected_content(self):
        rows = self.table.selectionModel().selectedRows()
        if not rows:
            return
        idx = rows[0].row()
        if idx >= len(self.results):
            return
        article = self.results[idx]
        text = (
            f"제목: {article['title']}\n"
            f"언론사: {article['press']}\n"
            f"날짜: {article['date']}\n"
            f"링크: {article['link']}\n\n"
            f"{article['content']}"
        )
        self.content_view.setPlainText(text)

    def save_csv(self):
        if not self.results:
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "CSV로 저장", "naver_news_result.csv", "CSV Files (*.csv)"
        )
        if not path:
            return
        save_to_csv(self.results, path)
        QMessageBox.information(self, "저장 완료", f"'{path}' 파일로 저장했습니다.")

    def save_excel(self):
        if not self.results:
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "엑셀로 저장", "naver_news_result.xlsx", "Excel Files (*.xlsx)"
        )
        if not path:
            return
        save_to_excel(self.results, path)
        QMessageBox.information(self, "저장 완료", f"'{path}' 파일로 저장했습니다.")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
